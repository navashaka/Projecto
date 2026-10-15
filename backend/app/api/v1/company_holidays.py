import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.api.deps import get_db
from app.schemas.company_holiday import (
    CompanyHolidayCreate,
    CompanyHolidayResponse,
)
from app.services.company_holiday_service import (
    create_company_holiday_service,
    delete_company_holiday_service,
    list_company_holidays_service,
)

router = APIRouter(
    prefix="/company-holidays",
    tags=["Company Holidays"],
)


@router.post("/", response_model=CompanyHolidayResponse)
def create_company_holiday(
    data: CompanyHolidayCreate,
    db: Session = Depends(get_db),
):
    return create_company_holiday_service(db, data)


@router.get("/", response_model=list[CompanyHolidayResponse])
def list_company_holidays(
    db: Session = Depends(get_db),
):
    return list_company_holidays_service(db)


@router.delete("/{holiday_id}")
def delete_company_holiday(
    holiday_id: int,
    db: Session = Depends(get_db),
):
    deleted = delete_company_holiday_service(db, holiday_id)

    if not deleted:
        raise HTTPException(
            status_code=404,
            detail="Company holiday not found",
        )

    return {"detail": "Company holiday deleted"}


@router.post("/upload")
async def upload_company_holidays(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Please upload an Excel file",
        )

    filename = file.filename.lower()

    if not filename.endswith((".xlsx", ".xls")):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx and .xls files are supported",
        )

    if filename.endswith(".xls"):
        raise HTTPException(
            status_code=400,
            detail="Please save the file as .xlsx and upload again",
        )

    file_content = await file.read()

    try:
        workbook = load_workbook(
            io.BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active

        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise HTTPException(
                status_code=400,
                detail="Excel file is empty",
            )

        column_names = [
            str(value).strip().lower() if value is not None else ""
            for value in headers
        ]

        required_columns = {"holiday_date", "holiday_name"}

        if not required_columns.issubset(set(column_names)):
            raise HTTPException(
                status_code=400,
                detail=(
                    "Excel must contain holiday_date and "
                    "holiday_name columns"
                ),
            )

        date_index = column_names.index("holiday_date")
        name_index = column_names.index("holiday_name")

        holidays = []

        for row_number, row in enumerate(rows, start=2):
            holiday_date = (
                row[date_index] if date_index < len(row) else None
            )
            holiday_name = (
                row[name_index] if name_index < len(row) else None
            )

            if holiday_date is None and holiday_name is None:
                continue

            if holiday_date is None or holiday_name is None:
                raise HTTPException(
                    status_code=400,
                    detail=f"Missing holiday date or name in row {row_number}",
                )

            if isinstance(holiday_date, datetime):
                holiday_date = holiday_date.date().isoformat()
            elif isinstance(holiday_date, date):
                holiday_date = holiday_date.isoformat()
            else:
                holiday_date = str(holiday_date).strip()

            holiday_name = str(holiday_name).strip()

            if not holiday_date or not holiday_name:
                raise HTTPException(
                    status_code=400,
                    detail=f"Missing holiday date or name in row {row_number}",
                )

            holidays.append(
                CompanyHolidayCreate(
                    holiday_date=holiday_date,
                    holiday_name=holiday_name,
                )
            )

        if not holidays:
            raise HTTPException(
                status_code=400,
                detail="No valid holiday records found in Excel",
            )

        uploaded_count = 0

        try:
            for holiday in holidays:
                create_company_holiday_service(db, holiday)
                uploaded_count += 1
        except Exception:
            db.rollback()
            raise

        return {
            "message": "Company holiday list uploaded successfully",
            "uploaded_count": uploaded_count,
        }

    except HTTPException:
        raise
    except Exception as error:
        db.rollback()
        raise HTTPException(
            status_code=400,
            detail=f"Unable to read Excel file: {error}",
        ) from error
    finally:
        if "workbook" in locals():
            workbook.close()
        await file.close()